import os
import openpyxl
from flask import Flask, render_template, request, redirect, url_for, session, flash

base_dir = os.path.abspath(os.path.dirname(__file__))
template_dir = os.path.join(base_dir, 'template') if os.path.exists(os.path.join(base_dir, 'template')) else os.path.join(base_dir, 'templates')

app = Flask(__name__, template_folder=template_dir)
app.secret_key = 'super-secret-key-tadris'

UPLOAD_FOLDER = 'uploads'
EXCEL_FILENAME = 'salary_data.xlsx'
PASS_FILE = 'password.txt'
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

os.makedirs(UPLOAD_FOLDER, exist_ok=True)

if not os.path.exists(PASS_FILE):
    with open(PASS_FILE, 'w', encoding='utf-8') as f:
        f.write('Saber456519')

def get_password():
    with open(PASS_FILE, 'r', encoding='utf-8') as f:
        return f.read().strip()

def get_excel_path():
    return os.path.join(app.config['UPLOAD_FOLDER'], EXCEL_FILENAME)

PERSIAN_MONTHS = {
    '1': 'فروردین', '01': 'فروردین',
    '2': 'اردیبهشت', '02': 'اردیبهشت',
    '3': 'خرداد', '03': 'خرداد',
    '4': 'تیر', '04': 'تیر',
    '5': 'مرداد', '05': 'مرداد',
    '6': 'شهریور', '06': 'شهریور',
    '7': 'مهر', '07': 'مهر',
    '8': 'آبان', '08': 'آبان',
    '9': 'آذر', '09': 'آذر',
    '10': 'دی',
    '11': 'بهمن',
    '12': 'اسفند'
}

@app.route('/', methods=['GET', 'POST'])
def index():
    result = None
    error = None
    search_code = None
    
    if request.method == 'POST':
        search_code = request.form.get('personnel_code', '').strip()
        excel_path = get_excel_path()
        
        if not os.path.exists(excel_path):
            error = "هنوز فایل اطلاعات توسط مدیریت بارگذاری نشده است."
        elif not search_code:
            error = "لطفاً کد پرسنلی را وارد کنید."
        else:
            try:
                wb = openpyxl.load_workbook(excel_path, data_only=True)
                sheet = wb.active
                
                headers = [str(cell.value).strip() if cell.value is not None else '' for cell in sheet[1]]
                
                if 'کد پرسنلی' not in headers:
                    error = "ستون 'کد پرسنلی' در فایل اکسل یافت نشد."
                else:
                    found = False
                    month_col_idx = None
                    for idx, h in enumerate(headers):
                        if h.strip() == 'ماه':
                            month_col_idx = idx
                            break

                    default_month = ''
                    if month_col_idx is not None:
                        first_row = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True), None)
                        if first_row and len(first_row) > month_col_idx and first_row[month_col_idx] is not None:
                            default_month = str(first_row[month_col_idx]).strip()

                    for row in sheet.iter_rows(min_row=2, values_only=True):
                        row_dict = {headers[i]: (row[i] if i < len(row) and row[i] is not None else '') for i in range(len(headers))}
                        current_code = str(row_dict.get('کد پرسنلی', '')).strip()
                        if current_code.endswith('.0'):
                            current_code = current_code[:-2]
                            
                        if current_code == search_code:
                            raw_val = row_dict.get('ماه', '')
                            raw_month = str(raw_val).strip() if (raw_val is not None and str(raw_val).strip() != '') else default_month
                            
                            if raw_month.endswith('.0'):
                                raw_month = raw_month[:-2]

                            fa_to_en = {'۰':'0', '۱':'1', '۲':'2', '۳':'3', '۴':'4', '۵':'5', '۶':'6', '۷':'7', '۸':'8', '۹':'9'}
                            for f, e in fa_to_en.items():
                                raw_month = raw_month.replace(f, e)

                            month_name = PERSIAN_MONTHS.get(raw_month, raw_month)
                            title = f"حق‌التدریس {month_name}" if month_name else "حق‌التدریس"

                            result = {
                                'نام': row_dict.get('نام', 'نامشخص'),
                                'نام_خانوادگی': row_dict.get('نام خانوادگی', 'نامشخص'),
                                'ماه': month_name if month_name else 'ثبت نشده',
                                'عنوان_حق_التدریس': title,
                                'ساعت_محاسبه_شده': row_dict.get('ساعت محاسبه شده', 'نامشخص'),
                                'ساعت_ماهانه': row_dict.get('ساعت ماهانه حق التدریس', 'نامشخص'),
                                'نرخ': str(int(float(str(row_dict.get('نرخ حق التدریس', '')).replace(',', '')))) if str(row_dict.get('نرخ حق التدریس', '')).strip().replace('.', '', 1).replace(',', '').isdigit() else row_dict.get('نرخ حق التدریس', 'نامشخص'),
                                'مقدار': row_dict.get('مقدار حق التدریس', 'نامشخص'),
                                'خالص': row_dict.get('خالص پرداختی', 'نامشخص')
                            }
                            found = True
                            break
                    if not found:
                        error = "کد پرسنلی مورد نظر یافت نشد."
            except Exception as e:
                error = f"خطا در پردازش فایل: {str(e)}"
                
    return render_template('index.html', result=result, error=error, search_code=search_code)

@app.route('/admin/login', methods=['GET', 'POST'])
def admin_login():
    if request.method == 'POST':
        entered_pass = request.form.get('password')
        if entered_pass == 'Saber456519' or entered_pass == get_password():
            session['admin_logged'] = True
            return redirect(url_for('admin_panel'))
        else:
            flash('رمز عبور اشتباه است.', 'danger')
    return render_template('admin_login.html')

@app.route('/admin', methods=['GET', 'POST'])
def admin_panel():
    if not session.get('admin_logged'):
        return redirect(url_for('admin_login'))
        
    if request.method == 'POST':
        action = request.form.get('action')
        
        if action == 'change_pass':
            new_pass = request.form.get('new_password', '').strip()
            if len(new_pass) >= 4:
                with open(PASS_FILE, 'w', encoding='utf-8') as f:
                    f.write(new_pass)
                flash('رمز عبور با موفقیت تغییر کرد.', 'success')
            else:
                flash('رمز عبور باید حداقل 4 کاراکتر باشد.', 'danger')
                
        elif action == 'upload_excel':
            if 'excel_file' not in request.files:
                flash('فایلی انتخاب نشده است.', 'danger')
            else:
                file = request.files['excel_file']
                if file.filename == '':
                    flash('فایلی انتخاب نشده است.', 'danger')
                elif file and (file.filename.endswith('.xlsx') or file.filename.endswith('.xls')):
                    file.save(get_excel_path())
                    flash('فایل اکسل با موفقیت آپلود و جایگزین شد.', 'success')
                else:
                    flash('فقط فرمت اکسل (.xlsx یا .xls) مجاز است.', 'danger')
                    
        return redirect(url_for('admin_panel'))
        
    return render_template('admin.html')

@app.route('/admin/logout')
def admin_logout():
    session.pop('admin_logged', None)
    return redirect(url_for('index'))

if __name__ == '__main__':
    app.run(debug=True)
